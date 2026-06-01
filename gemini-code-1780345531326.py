import io
import re

import streamlit as st

try:
    import docx
except ImportError:
    docx = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    import fitz
except ImportError:
    fitz = None

ALLOWED_EXTENSIONS = ["pdf", "docx", "pptx", "txt", "png", "jpg", "jpeg"]

st.set_page_config(
    page_title="Push Your Limits",
    page_icon="🩺",
    layout="wide",
)

st.markdown(
    """
    <style>
    .main { background-color: #f4f6f4; }
    h1, h2, h3 { color: #2c4a3e; font-family: Georgia, serif; }
    .stButton>button {
        background-color: #3d6453;
        color: white;
        border-radius: 8px;
        border: none;
        padding: 0.75rem 1rem;
    }
    .stButton>button:hover { background-color: #2c4a3e; color: white; }
    .card { background: white; padding: 1.5rem; border-radius: 18px; box-shadow: 0 14px 30px rgba(0,0,0,0.06); margin-bottom: 1.5rem; }
    .success-text { color: #2e7d32; font-weight: 700; }
    </style>
    """,
    unsafe_allow_html=True,
)

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False
if "uploads" not in st.session_state:
    st.session_state.uploads = []
if "current_upload" not in st.session_state:
    st.session_state.current_upload = None
if "quiz_index" not in st.session_state:
    st.session_state.quiz_index = 0
if "quiz_score" not in st.session_state:
    st.session_state.quiz_score = 0
if "quiz_mode" not in st.session_state:
    st.session_state.quiz_mode = False


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_pdf(file_bytes):
    if not PdfReader:
        return ""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        text = []
        for page in reader.pages:
            text.append(page.extract_text() or "")
        return "\n".join(text).strip()
    except Exception:
        return ""


def extract_text_from_docx(file_bytes):
    if not docx:
        return ""
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        return "\n".join([p.text for p in document.paragraphs if p.text.strip()])
    except Exception:
        return ""


def extract_text_from_pptx(file_bytes):
    if not Presentation:
        return ""
    try:
        presentation = Presentation(io.BytesIO(file_bytes))
        lines = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(shape.text.strip())
        return "\n".join(lines)
    except Exception:
        return ""


def extract_text_from_txt(file_bytes):
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def extract_images_from_pdf(file_bytes):
    images = []
    if not fitz:
        return images
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page_index in range(len(doc)):
            for img_info in doc.get_page_images(page_index):
                xref = img_info[0]
                image = doc.extract_image(xref)
                images.append(image.get("image"))
    except Exception:
        pass
    return images


def summarize_text(text, limit=280):
    sentences = re.split(r"(?<=[.!?])\s+", text)
    summary = []
    for sentence in sentences:
        if sentence.strip():
            summary.append(sentence.strip())
            if len(summary) >= 3:
                break
    return " ".join(summary)[:limit]


def generate_quiz_items(text):
    questions = []
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if len(s.strip()) > 30]
    for index, sentence in enumerate(sentences[:5], start=1):
        words = re.findall(r"\b[A-Za-z]{4,}\b", sentence)
        if len(words) < 3:
            continue
        answer = words[0]
        distractors = [w for w in words[1:5] if w.lower() != answer.lower()][:3]
        if len(distractors) < 3:
            distractors = [f"Option {i}" for i in range(1, 4)]
        choices = [answer] + distractors
        questions.append({
            "question": f"Which term best matches this concept?\n\n{sentence[:120].rstrip()}...",
            "choices": choices,
            "answer": answer,
        })
    return questions


def process_uploaded_file(uploaded_file):
    file_bytes = uploaded_file.read()
    extension = uploaded_file.name.rsplit(".", 1)[-1].lower()
    content = ""
    images = []

    if extension == "pdf":
        content = extract_text_from_pdf(file_bytes)
        images = extract_images_from_pdf(file_bytes)
    elif extension == "docx":
        content = extract_text_from_docx(file_bytes)
    elif extension == "pptx":
        content = extract_text_from_pptx(file_bytes)
    elif extension == "txt":
        content = extract_text_from_txt(file_bytes)
    elif extension in ["png", "jpg", "jpeg"]:
        images = [file_bytes]

    quiz = generate_quiz_items(content)
    summary = summarize_text(content or "Upload a supported study document to generate questions.")
    return {
        "name": uploaded_file.name,
        "content": content,
        "summary": summary,
        "quiz": quiz,
        "images": images,
    }

st.title("🩺 Push Your Limits")
st.write("A calm medical study portal. Upload your own study files and generate quizzes from your material.")
st.markdown("---")

with st.sidebar:
    st.header("Editor Access")
    if not st.session_state.authenticated:
        with st.form("login_form"):
            username = st.text_input("Username")
            password = st.text_input("Password", type="password")
            submitted = st.form_submit_button("Login")
            if submitted:
                if username == "admin" and password == "study2026":
                    st.session_state.authenticated = True
                    st.success("Editor access granted.")
                else:
                    st.error("Invalid username or password.")
        st.markdown("---")
        st.write("Need upload access? Send a request to the admin.")
        if st.button("Request editor access"):
            st.markdown(
                "[Open email client](mailto:luffy5656taro@gmail.com?subject=Editor+Access+Request&body=Please+grant+me+editor+access.)"
            )
    else:
        st.success("Logged in as editor")
        if st.button("Logout"):
            st.session_state.authenticated = False
            st.experimental_rerun()

if st.session_state.authenticated:
    st.subheader("Upload your own study material")
    uploaded_file = st.file_uploader(
        "Choose a study file to upload:",
        type=["pdf", "docx", "pptx", "txt", "png", "jpg", "jpeg"],
    )
    if uploaded_file:
        if allowed_file(uploaded_file.name):
            uploaded_data = process_uploaded_file(uploaded_file)
            st.session_state.uploads.append(uploaded_data)
            st.success(f"Uploaded {uploaded_file.name} successfully.")
        else:
            st.error("This file type is not supported.")

    if st.session_state.uploads:
        st.markdown("---")
        st.subheader("Uploaded files")
        selected_name = st.selectbox("Choose an uploaded file", [u["name"] for u in st.session_state.uploads])
        current = next(u for u in st.session_state.uploads if u["name"] == selected_name)
        st.write(f"**Summary:** {current['summary']}")
        if current["images"]:
            st.image(current["images"][0], caption="Extracted image from file", use_column_width=True)
        if current["quiz"]:
            if st.button("Start quiz for this file"):
                st.session_state.current_upload = current
                st.session_state.quiz_index = 0
                st.session_state.quiz_score = 0
                st.session_state.quiz_mode = True

    if st.session_state.quiz_mode and st.session_state.current_upload:
        quiz = st.session_state.current_upload["quiz"]
        if quiz and st.session_state.quiz_index < len(quiz):
            question = quiz[st.session_state.quiz_index]
            st.markdown(f"### Question {st.session_state.quiz_index + 1} of {len(quiz)}")
            st.markdown(question["question"])
            choice = st.radio("Select the best answer", question["choices"], key=st.session_state.quiz_index)
            if st.button("Submit answer"):
                if choice == question["answer"]:
                    st.success("Correct! Great work.")
                    st.session_state.quiz_score += 1
                else:
                    st.error(f"Incorrect. The correct answer is: {question['answer']}")
                st.session_state.quiz_index += 1
                st.experimental_rerun()
        else:
            st.balloons()
            st.success("Quiz complete!")
            total = len(st.session_state.current_upload["quiz"])
            score = st.session_state.quiz_score
            st.write(f"You answered {score} out of {total} correctly.")
            st.write(f"Result: {int(score / total * 100)}%")
            if st.button("Restart quiz"):
                st.session_state.quiz_mode = False
                st.session_state.current_upload = None
                st.experimental_rerun()
else:
    st.warning("Editor login required to upload files and generate quizzes.")
    st.write("Please use the sidebar login to access upload features.")


    
        
        


            


