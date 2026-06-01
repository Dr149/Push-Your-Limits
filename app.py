import json
import os
import re
import secrets
import smtplib
import uuid
from email.message import EmailMessage
from pathlib import Path

from dotenv import load_dotenv
from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
import random

try:
    import fitz
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_FOLDER = BASE_DIR / "uploads"
DATA_FOLDER = BASE_DIR / "data"
IMAGE_FOLDER = UPLOAD_FOLDER / "images"
USER_FILE = DATA_FOLDER / "users.json"
CONTENT_FILE = DATA_FOLDER / "content.json"
ALLOWED_EXTENSIONS = {"txt", "pdf", "docx", "pptx", "png", "jpg", "jpeg"}
EMAIL_RECIPIENT = "luffy5656taro@gmail.com"

load_dotenv()
def ensure_json(path, default):
    if not path.exists() or path.stat().st_size == 0:
        with path.open("w", encoding="utf-8") as fh:
            json.dump(default, fh, indent=2)


def generate_unique_filename(filename):
    safe_name = secure_filename(filename)
    unique_token = uuid.uuid4().hex
    return f"{Path(safe_name).stem}_{unique_token}{Path(safe_name).suffix}"


app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", secrets.token_hex(16))
app.config["UPLOAD_FOLDER"] = str(UPLOAD_FOLDER)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024

for folder in (UPLOAD_FOLDER, IMAGE_FOLDER, DATA_FOLDER):
    folder.mkdir(parents=True, exist_ok=True)

ensure_json(USER_FILE, [{
    "username": "admin",
    "password": generate_password_hash("study2026"),
    "is_editor": True,
}])
ensure_json(CONTENT_FILE, {"files": [], "videos": [], "quizzes": []})


def load_json(path):
    try:
        with path.open("r", encoding="utf-8") as fh:
            return json.load(fh)
    except (json.JSONDecodeError, FileNotFoundError):
        if path == USER_FILE:
            default = []
        else:
            default = {"files": [], "videos": [], "quizzes": []}
        save_json(path, default)
        return default


def save_json(path, data):
    with path.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2)


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_pdf(path):
    if not PdfReader:
        return ""
    try:
        reader = PdfReader(str(path))
        text = []
        for page in reader.pages:
            text.append(page.extract_text() or "")
        return "\n".join(text).strip()
    except Exception:
        return ""


def extract_text_from_docx(path):
    if not docx:
        return ""
    try:
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except Exception:
        return ""


def extract_text_from_pptx(path):
    if not Presentation:
        return ""
    try:
        presentation = Presentation(str(path))
        text_parts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text.strip())
        return "\n".join(text_parts)
    except Exception:
        return ""


def extract_images_from_docx(path, file_id):
    images = []
    if not docx:
        return images
    try:
        document = docx.Document(str(path))
        for i, rel in enumerate(document.part.rels.values(), start=1):
            if "image" in rel.target_ref:
                image_bytes = rel.target_part.blob
                ext = Path(rel.target_ref).suffix or ".png"
                filename = f"file_{file_id}_docx_image_{i}{ext}"
                target_path = IMAGE_FOLDER / filename
                target_path.write_bytes(image_bytes)
                images.append(filename)
    except Exception:
        pass
    return images


def extract_images_from_pdf(path, file_id):
    images = []
    if not fitz:
        return images
    try:
        doc = fitz.open(str(path))
        for page_index in range(len(doc)):
            for img_index, img in enumerate(doc.get_page_images(page_index), start=1):
                xref = img[0]
                image = doc.extract_image(xref)
                ext = image.get("ext", "png")
                filename = f"file_{file_id}_pdf_image_{page_index + 1}_{img_index}.{ext}"
                target_path = IMAGE_FOLDER / filename
                target_path.write_bytes(image["image"])
                images.append(filename)
    except Exception:
        pass
    return images


def summarize_text(text, limit=280):
    sentences = re.split(r"(?<=[.!?])\\s+", text)
    summary = []
    count = 0
    for sentence in sentences:
        if sentence.strip():
            summary.append(sentence.strip())
            count += 1
            if count >= 3:
                break
    return " ".join(summary)[:limit]


def generate_quiz_items(text, images):
    questions = []
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\\s+", text) if len(s.strip()) > 30]
    if not sentences:
        return questions
    for index, sentence in enumerate(sentences[:5], start=1):
        words = re.findall(r"\\b[A-Za-z]{4,}\\b", sentence)
        if len(words) < 3:
            continue
        sorted_words = sorted(set(words), key=lambda item: -len(item))
        answer = sorted_words[0]
        distractors = [w for w in sorted_words[1:5] if w.lower() != answer.lower()][:3]
        if len(distractors) < 3:
            distractors = [f"Option {i}" for i in range(1, 4)]
        choices = [answer] + distractors
        random.shuffle(choices)
        questions.append(
            {
                "id": index,
                "question": f"Which term best matches this concept?\n\n{sentence[:120].rstrip()}...",
                "choices": choices,
                "answer": answer,
                "image": images[index - 1] if index - 1 < len(images) else None,
            }
        )
    return questions


@app.context_processor
def inject_brand():
    return {"brand": "Push Your Limits"}


def send_access_request(name, role, email=None):
    smtp_host = os.environ.get("SMTP_HOST")
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER")
    smtp_password = os.environ.get("SMTP_PASSWORD")
    subject = "Push Your Limits Editor Access Request"
    body = f"Name: {name}\nRequested role: {role}\nUser email: {email or 'not provided'}\n"
    if smtp_host and smtp_user and smtp_password:
        try:
            message = EmailMessage()
            message["From"] = smtp_user
            message["To"] = EMAIL_RECIPIENT
            message["Subject"] = subject
            message.set_content(body)
            with smtplib.SMTP(smtp_host, smtp_port) as server:
                server.starttls()
                server.login(smtp_user, smtp_password)
                server.send_message(message)
            return True
        except Exception:
            return False
    print("Email request:\n", body)
    return False


def current_user():
    username = session.get("username")
    if not username:
        return None
    users = load_json(USER_FILE)
    return next((u for u in users if u["username"] == username), None)


def require_editor():
    user = current_user()
    return user and user.get("is_editor")


@app.route("/")
def home():
    content = load_json(CONTENT_FILE)
    files = [f for f in content["files"] if not f.get("hidden")]
    videos = [v for v in content["videos"] if not v.get("hidden")]
    quizzes = [q for q in content["quizzes"] if not q.get("hidden")]
    return render_template("index.html", files=files, videos=videos, quizzes=quizzes)


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"].strip()
        users = load_json(USER_FILE)
        user = next((u for u in users if u["username"] == username), None)
        if user and check_password_hash(user["password"], password):
            session["username"] = username
            flash("Welcome back, editor.", "success")
            return redirect(url_for("editor"))
        flash("Credentials not recognized. Please try again.", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.pop("username", None)
    flash("You have been logged out.", "info")
    return redirect(url_for("home"))


@app.route("/request-access", methods=["GET", "POST"])
def request_access():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        role = request.form.get("role", "viewer")
        email = request.form.get("email", "").strip()
        if not name:
            flash("Please enter your full name.", "warning")
            return redirect(url_for("request_access"))
        sent = send_access_request(name, role, email)
        if sent:
            flash("Access request sent successfully.", "success")
        else:
            flash("Your request was recorded, and the admin will review it.", "info")
        return redirect(url_for("home"))
    return render_template("request_access.html")


@app.route("/editor")
def editor():
    if not require_editor():
        flash("Editor access required.", "warning")
        return redirect(url_for("login"))
    content = load_json(CONTENT_FILE)
    return render_template("editor.html", content=content)


@app.route("/editor/upload", methods=["POST"])
def upload():
    if not require_editor():
        return redirect(url_for("login"))
    uploaded_file = request.files.get("file")
    title = request.form.get("title", "Untitled").strip()
    if not uploaded_file or not allowed_file(uploaded_file.filename):
        flash("Please upload a supported file type.", "danger")
        return redirect(url_for("editor"))
    filename = generate_unique_filename(uploaded_file.filename)
    destination = UPLOAD_FOLDER / filename
    uploaded_file.save(destination)

    content = load_json(CONTENT_FILE)
    file_id = (content["files"] and max(item["id"] for item in content["files"])) + 1 if content["files"] else 1
    text = ""
    images = []
    ext = destination.suffix.lower()
    if ext == ".txt":
        text = destination.read_text(encoding="utf-8", errors="ignore")
    elif ext == ".pdf":
        text = extract_text_from_pdf(destination)
        images = extract_images_from_pdf(destination, file_id)
    elif ext == ".docx":
        text = extract_text_from_docx(destination)
        images = extract_images_from_docx(destination, file_id)
    elif ext == ".pptx":
        text = extract_text_from_pptx(destination)
    elif ext in {".png", ".jpg", ".jpeg"}:
        images = [filename]

    quiz = generate_quiz_items(text, images)
    content["files"].append(
        {
            "id": file_id,
            "title": title or filename,
            "filename": filename,
            "type": ext.lstrip("."),
            "hidden": False,
            "summary": summarize_text(text or "No text available."),
            "images": images,
            "quiz_id": file_id,
        }
    )
    if quiz:
        content["quizzes"].append(
            {
                "id": file_id,
                "title": f"Quiz from {title or filename}",
                "hidden": False,
                "questions": quiz,
            }
        )
    save_json(CONTENT_FILE, content)
    flash("File uploaded and quiz generated successfully.", "success")
    return redirect(url_for("editor"))


@app.route("/editor/toggle/<string:section>/<int:item_id>")
def toggle_visibility(section, item_id):
    if not require_editor():
        return redirect(url_for("login"))
    content = load_json(CONTENT_FILE)
    items = content.get(section, [])
    item = next((x for x in items if x["id"] == item_id), None)
    if item:
        item["hidden"] = not item.get("hidden", False)
        save_json(CONTENT_FILE, content)
        flash(f"Visibility updated for {section} item.", "success")
    return redirect(url_for("editor"))


@app.route("/editor/add-video", methods=["POST"])
def add_video():
    if not require_editor():
        return redirect(url_for("login"))
    title = request.form.get("title", "New video").strip()
    url = request.form.get("url", "").strip()
    if not url:
        flash("Please provide a video link.", "warning")
        return redirect(url_for("editor"))
    content = load_json(CONTENT_FILE)
    video_id = (content["videos"] and max(v["id"] for v in content["videos"])) + 1 if content["videos"] else 1
    content["videos"].append(
        {"id": video_id, "title": title, "url": url, "hidden": False}
    )
    save_json(CONTENT_FILE, content)
    flash("Video added successfully.", "success")
    return redirect(url_for("editor"))


@app.route("/quiz/<int:quiz_id>")
def quiz(quiz_id):
    content = load_json(CONTENT_FILE)
    quiz_data = next((q for q in content["quizzes"] if q["id"] == quiz_id and not q.get("hidden")), None)
    if not quiz_data:
        flash("Quiz not found.", "warning")
        return redirect(url_for("home"))
    return render_template("quiz.html", quiz=quiz_data)


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(str(UPLOAD_FOLDER), filename)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
